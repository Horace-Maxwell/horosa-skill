"""完整排盘报告的版式与信息质量（v0.43.1-dev，2026-10-24 独立评审后）：盘图、表格、行首禁则、补字链、人读附录。"""
from __future__ import annotations

import re
from pathlib import Path

import pytest

from horosa_skill.reports import builder, figures, renderers

# ── 盘图：只取引擎给的数据，不算 ─────────────────────────────────────────


def _wheel_chart(asc: float = 85.0) -> dict:
    houses = [{"id": f"House{i}", "lon": (asc + (i - 1) * 30) % 360} for i in range(1, 13)]
    names = ["Sun", "Moon", "Mercury", "Venus", "Mars", "Jupiter", "Saturn", "Uranus", "Neptune", "Pluto", "North Node", "South Node"]
    objects = [{"id": n, "lon": (i * 27.3) % 360, "lonspeed": -0.1 if n == "Mercury" else 0.5} for i, n in enumerate(names)]
    objects += [{"id": "Asc", "lon": asc}, {"id": "MC", "lon": (asc + 270) % 360}]
    return {"houses": houses, "objects": objects}


def test_wheel_uses_engine_positions_and_engine_aspects_only() -> None:
    data = {"chart": _wheel_chart(), "aspects": {"normalAsp": {
        "Sun": {"Applicative": [{"id": "Moon", "asp": 90}, {"id": "Mars", "asp": 0}, {"id": "Pars Fortuna", "asp": 120}]},
        "Moon": {"Separative": [{"id": "Sun", "asp": 90}]},
    }}}
    wheel, = figures.extract_figures("chart", data, "星盘")
    assert wheel["kind"] == "wheel" and wheel["title"] == "星盘"  # 不再叠成「星盘盘」
    assert len(wheel["cusps"]) == 12 and wheel["axes"][0]["lon"] == 85.0
    mercury = next(b for b in wheel["bodies"] if b["id"] == "Mercury")
    assert mercury["retro"] is True and mercury["label"] == "水"
    assert not next(b for b in wheel["bodies"] if b["id"] == "North Node")["retro"]
    chart = _wheel_chart()
    next(o for o in chart["objects"] if o["id"] == "North Node")["lonspeed"] = -0.05
    nodes, = figures.extract_figures("chart", {"chart": chart}, "星盘")
    assert next(b for b in nodes["bodies"] if b["id"] == "North Node")["retro"]  # 交点逆行照引擎速度标，与位置表一致
    # 只画引擎列出的相位：合相（0）不连线，虚点不在盘上；同一对只画一次。
    assert wheel["aspects"] == [{"a": "Sun", "b": "Moon", "asp": 90}]


def test_wheel_sign_names_match_the_engine_text_and_babylon_has_no_houses_or_aspects() -> None:
    wheel, = figures.extract_figures("chart", {"chart": _wheel_chart()}, "星盘")
    assert [s["label"] for s in wheel["signs"]][::5] == ["牡羊", "室女", "宝瓶"]  # 与导出正文同一套译名
    aspects = {"Sun": {"Applicative": [{"id": "Moon", "asp": 90}]}}
    babylon, = figures.extract_figures("babylon", {"chart": _wheel_chart(), "aspects": {"normalAsp": aspects}}, "巴比伦占星")
    # 引擎起盘信息：「本体系无十二宫位、无相位、无上升点」——盘图不能画出来。
    assert babylon["houses"] is False and babylon["aspects"] == []
    horary, = figures.extract_figures("horary", {"chart": _wheel_chart(), "raw": {"aspects": {"normalAsp": aspects}}}, "卜卦盘")
    assert horary["aspects"] == [{"a": "Sun", "b": "Moon", "asp": 90}]  # 卜卦 / 择日的相位在 raw 里


def test_hexagram_draws_the_changed_gua_only_when_it_agrees_with_the_moving_lines() -> None:
    lines = [{"value": 1, "change": i in (2, 5), "god": g} for i, g in enumerate(["青龙", "朱雀", "勾陈", "螣蛇", "白虎", "玄武"])]
    descriptions = {"111111": {"desc": "乾为天", "yao": [1] * 6}, "110110": {"desc": "兑为泽", "yao": [1, 1, 0, 1, 1, 0]},
                    "101010": {"desc": "水火既济", "yao": [1, 0, 1, 0, 1, 0]}}
    ok, = figures.extract_figures("sixyao", {"current_code": "111111", "changed_code": "110110", "lines": lines,
                                             "descriptions": descriptions})
    assert [c["title"] for c in ok["columns"]] == ["本卦 · 乾为天", "变卦 · 兑为泽"]
    assert ok["columns"][0]["moving"] == [False, False, True, False, False, True]
    # 六神不画：data.lines 的六神可以与引擎「断卦结构」表（按日干起）不一致，图与表打架时读者无所适从。
    assert ok["columns"][0]["lines"] == [""] * 6 and "断卦结构" in ok["text"]
    # 调用方给的 changed_code 与动爻对不上（样例载荷正是如此）：只画本卦，之卦以正文为准，不自己翻爻去推。
    mismatch, = figures.extract_figures("sixyao", {"current_code": "111111", "changed_code": "101010", "lines": lines,
                                                   "descriptions": descriptions})
    assert [c["title"] for c in mismatch["columns"]] == ["本卦 · 乾为天"] and "之卦见下文" in mismatch["text"]


def test_wheel_titles_and_tradition_exclusions() -> None:
    returned = figures.extract_figures("solarreturn", {"dirChart": {"chart": _wheel_chart()}, "natalChart": {"chart": _wheel_chart(10)}},
                                       "推运盘-太阳返照")
    assert [w["title"] for w in returned] == ["太阳返照盘", "本命盘"]
    assert figures._chart_title("量化盘（汉堡学派中点）") == "量化盘（汉堡学派中点）" and figures._chart_title("星运-Balbillus") == "Balbillus 盘"
    pair = figures.extract_figures("relative", {"inner": {"chart": _wheel_chart()}, "outer": {"chart": _wheel_chart(10)}}, "关系盘")
    assert [w["title"] for w in pair] == ["甲方星盘", "乙方星盘"]
    seasons = figures.extract_figures("jieqi_year", {"charts": {k: {"chart": _wheel_chart()} for k in ("夏至", "冬至", "春分", "秋分")}}, "节气盘")
    assert [w["title"] for w in seasons] == ["春分盘", "夏至盘", "秋分盘", "冬至盘"]
    # 七政四余 / 印度 / 宿占不套西洋轮（宫制、星名、盘式都不同）。
    assert figures.extract_figures("guolao_chart", {"chart": _wheel_chart()}, "七政四余") == []
    assert figures.extract_figures("india_chart", {"chart": _wheel_chart()}, "印度律盘") == []
    # 缺宫的盘宁可不画。
    broken = _wheel_chart()
    broken["houses"] = broken["houses"][:11]
    assert figures.extract_figures("chart", {"chart": broken}, "星盘") == []


def test_label_spreading_keeps_order_and_gap() -> None:
    lons = [6.7, 10.6, 13.4, 16.8, 29.1, 61.9, 67.1, 118.4, 150.0, 169.6, 298.4, 308.6]
    shown = figures._spread(lons, 10)
    # 相邻显示角至少 10°、先后次序不变（此前两两推挤会把「火」推到「日」后面，标签叠在一起）。
    unwrapped = [shown[0]] + [shown[0] + (s - shown[0]) % 360 for s in shown[1:]]
    assert unwrapped == sorted(unwrapped)
    assert all((shown[(i + 1) % 12] - shown[i]) % 360 >= 10 - 1e-6 for i in range(12))
    assert figures._spread([355, 357, 2, 4], 10) == pytest.approx([344.5, 354.5, 4.5, 14.5])


def test_ziwei_qimen_liureng_grids_are_placed_by_branch_and_trigram() -> None:
    branches = "子丑寅卯辰巳午未申酉戌亥"
    houses = [{"name": f"宫{i}", "ganzi": f"甲{b}", "starsMain": [{"name": "紫微", "starlight": "庙", "sihua": "禄"}] if i == 0 else [],
               "isLife": i == 3, "isBody": i == 5, "direction": [2 + i * 10, 11 + i * 10]} for i, b in enumerate(branches)]
    ziwei, = figures.extract_figures("ziwei_birth", {"chart": {"houses": houses, "wuxingJuText": "水二局", "lifeMaster": "廉贞",
                                                              "bodyMaster": "天梁"}})
    by_branch = {c["branch"]: c for c in ziwei["cells"]}
    assert by_branch["子"]["lines"][0] == "紫微庙·禄" and by_branch["丑"]["lines"][0] == "（空宫）"
    assert by_branch["卯"]["emphasis"] == "main" and by_branch["巳"]["title"].endswith("（身）")
    assert "水二局" in ziwei["center"]["lines"]

    order = "巽离坤震中兑艮坎乾"
    cells = [{"palaceName": g, "god": "符" if g == "坤" else "地", "tianXing": "心", "door": "杜", "tianGan": "戊", "diGan": "丙",
              "isZhiFu": g == "坤", "isCenter": g == "中", "hasKongWang": g == "坎"} for g in order]
    qimen, = figures.extract_figures("qimen", {"pan": {"cells": cells, "juText": "阳遁六局"}})
    kun = next(c for c in qimen["cells"] if c["gua"] == "坤")
    assert kun["title"] == "坤二宫" and kun["lines"][:3] == ["值符", "天心　戊", "杜门　丙"]
    assert "值符" not in kun["label"]  # 八神值符已写在宫里，不再重复标
    assert next(c for c in qimen["cells"] if c["gua"] == "坎")["label"] == "空亡"

    layout = {"downZi": list(branches), "upZi": list(branches[5:] + branches[:5]), "houseTianJiang": ["贵人"] * 12}
    liureng, = figures.extract_figures("liureng_gods", {"headless_liureng": {
        "layout": layout, "ke": {"raw": [["天后", "卯", "辛"], ["天空", "申", "卯"], ["贵人", "寅", "酉"], ["白虎", "未", "寅"]]},
        "sanChuan": {"cuang": ["己未", "空子", "丁巳"], "tianJiang": ["白虎", "朱雀", "玄武"], "liuQin": ["父母", "子孙", "官鬼"],
                     "name": "涉害课"}}})
    assert liureng["center"]["title"] == "涉害课" and "初传 己未 白虎 父母" in liureng["center"]["lines"]
    assert {c["branch"]: c["label"] for c in liureng["cells"]}["子"] == "巳"


def test_figure_text_is_localized_by_key_whitelist() -> None:
    from horosa_skill.textscript import TRADITIONAL
    from horosa_skill.textscript.localize import localize_document

    figure = {"kind": "grid9", "title": "奇门盘", "cells": [{"gua": "离", "title": "离九宫", "lines": ["杜门　丙"], "label": "马星"}]}
    out = localize_document({"figures": [figure]}, TRADITIONAL)["figures"][0]
    assert out["title"] == "奇門盤" and out["cells"][0]["lines"] == ["杜門　丙"] and out["cells"][0]["label"] == "馬星"
    assert out["cells"][0]["gua"] == "离"  # 位置键不转（渲染器按它定格）


# ── 表格与文字 ──────────────────────────────────────────────────────────


def test_wide_short_tables_are_transposed_and_empty_columns_dropped() -> None:
    header = ["柱", "干支", "藏干", "十神", "纳音", "长生", "星运", "自坐", "空亡"]
    rows = [["年柱", "戊申", "庚劫", "印", "大驿土", "长生", "临官", "病", "寅卯"], ["日柱", "辛酉", "辛比", "日元", "石榴木", "胎", "帝旺", "帝旺", "子丑"]]
    t_header, t_rows = renderers._transpose_wide_table(header, rows)
    assert t_header == ["柱", "年柱", "日柱"] and t_rows[0] == ["干支", "戊申", "辛酉"] and len(t_rows) == 8
    assert renderers._transpose_wide_table(header[:5], [r[:5] for r in rows]) == (header[:5], [r[:5] for r in rows])

    h, r = renderers._drop_empty_columns(["点", "位置", "逆行"], [["日", "16°", "—"], ["月", "29°", ""], ["水", "10°", "—"]])
    assert h == ["点", "位置"] and r[0] == ["日", "16°"]
    assert renderers._drop_empty_columns(["点", "逆行", "x"], [["日", "逆行", "1"]] * 3)[0] == ["点", "逆行", "x"]


def test_header_text_gets_room_and_long_floats_are_rounded_for_humans() -> None:
    widths = renderers._column_widths(["板块", "起始年", "起始年龄", "大运", "流年"], [["板块1", "2037", "10岁", "丁巳", "甲" * 40]], 470, 30)
    assert widths[1] >= len("起始年") * 2 * 4.6 + 11 - 0.01  # 表头放得下，不折成「起始 / 年」
    squeezed = renderers._column_widths(["起始年龄", "流年"], [["1", "甲" * 60]], 120, 30)
    assert squeezed[0] >= len("起始年龄") * 2 * 4.6 + 11 - 0.01 and abs(sum(squeezed) - 120) < 0.01  # 挤的时候也不跌破表头宽
    assert renderers._display_number("误差0.7548645849180389，度数 16.5") == "误差0.7549，度数 16.5"
    assert "0.7549" in renderers._md_inline("误差0.7548645849180389") and "0.75486" not in renderers._md_inline("误差0.7548645849180389")
    assert renderers._display_number("v0.43.0 · 2028-04-06") == "v0.43.0 · 2028-04-06"


def test_kinsoku_and_time_colons() -> None:
    from reportlab.lib import textsplit

    renderers._extend_kinsoku()
    assert all(ch in textsplit.ALL_CANNOT_START for ch in "，：；）、。")
    before = textsplit.ALL_CANNOT_START
    renderers._extend_kinsoku()
    assert textsplit.ALL_CANNOT_START == before  # 幂等
    assert renderers._is_time_colon("2028-04-06 09", "33:00 周四")
    assert not renderers._is_time_colon("日期", "2028-04-06 09:33:00")


def test_glyph_chain_picks_the_first_font_that_has_each_char() -> None:
    chain = [("ExtB", {ord("𥘩")}), ("Sym", {ord("䷀"), ord("♃"), ord("ḥ")})]
    out = renderers._wrap_missing_glyphs_chain("乾䷀♃ ḥ𥘩a", {ord("乾"), ord("a"), ord(" ")}, chain)
    assert out == '乾<font name="Sym">䷀♃</font> <font name="Sym">ḥ</font><font name="ExtB">𥘩</font>a'
    # 拉丁扩展（ḥ ṭ ṇ）此前因阈值 0x2E80 永远不补；ASCII 仍然永不被包。
    assert renderers._wrap_missing_glyphs_chain("<b>&amp;</b>", set(), [("X", {ord(c) for c in "<b>&amp;/"})]) == "<b>&amp;</b>"


def test_long_listings_and_narrow_tables_flow_in_columns_with_reading_order(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    lines = [f"Venus | Mars = {i}˚Aries" for i in range(160)]
    rows = [[str(2028 + i), str(i + 1), "甲子", "乙丑"] for i in range(99)]
    document = {"title": "分栏", "style": "chart", "sections": [
        {"id": "a", "title": "中点", "body": "\n".join(lines), "human_visible": True, "group": "chart"},
        {"id": "b", "title": "小运", "body": "| 年份 | 虚岁 | 小运 | 流年 |\n| --- | --- | --- | --- |\n" + "\n".join(
            "| " + " | ".join(r) + " |" for r in rows), "human_visible": True, "group": "chart"},
    ]}
    from reportlab.platypus import SimpleDocTemplate

    stories: list[list] = []
    original = SimpleDocTemplate.build

    def spy(doc, flowables, *args, **kwargs):
        stories.append(list(flowables))
        return original(doc, flowables, *args, **kwargs)

    monkeypatch.setattr(SimpleDocTemplate, "build", spy)
    out = renderers.render_report(document, output_path=tmp_path / "p.pdf", format_name="pdf")
    assert out["format"] == "pdf"
    panels = [f for f in stories[-1] if type(f).__name__ == "Panels"]
    # 两段都分栏排，且一行不丢、顺序不乱（分页时按剩余高度现切，见 renderers.Panels）。
    assert [len(p.rows) for p in panels] == [160, 99] and panels[0].panels >= 2 and panels[1].panels >= 2
    assert panels[0].rows[0] == [lines[0]] and panels[1].rows[-1] == rows[-1]
    pages = len(re.findall(rb"/Type /Page[^s]", (tmp_path / "p.pdf").read_bytes()))
    assert 0 < pages <= 6  # 一行一段时这份要 9 页以上


# ── 人读附录与标题 ──────────────────────────────────────────────────────


def test_settings_and_basis_read_like_text_not_machine_values() -> None:
    assert builder._setting_text("tradition", False) == "否" and builder._setting_text("predictive", True) == "是"
    assert builder._setting_text("after23NewDay", 1) == "是" and builder._setting_text("timeAlg", 0) == "真太阳时"
    # 宫制 / 黄道 / 性别不在附录重复：技法卡是请求口径，巴比伦（恒星黄道、无宫）卡上却是回归黄道 / Alcabitus；实际值看正文。
    assert builder._setting_text("hsys", 1) is None and builder._setting_text("zodiacal", 0) is None
    assert builder._setting_text("gender", True) is None
    # 没有人读码表的枚举不印（贵人法 2、流派 donggong、主限法 core_alchabitius）。
    assert builder._setting_text("guirengType", 2) is None and builder._setting_text("pdMethod", "core_alchabitius") is None
    assert builder._setting_text("school", "通行版") == "通行版"
    assert builder._coord_text("121e28") == "东经 121°28′" and builder._coord_text("31n13") == "北纬 31°13′"
    assert builder._coord_text("Shanghai") == "Shanghai"
    basis = builder.ReportBuilder()._chart_basis_section({"gender": True, "name": ["111111", "000000"], "lat": "31n13"},
                                                         {"technique_label": "紫微斗数"})
    assert basis["items"]["性别"] == "男" and "姓名" not in basis["items"] and basis["items"]["纬度"] == "北纬 31°13′"


def test_report_titles_drop_ui_tab_names_and_group_prefixes() -> None:
    assert builder._report_technique_name("nongli_time", "其他页面") == "农历时间换算"
    assert builder._report_technique_name("solarreturn", "推运盘-太阳返照") == "太阳返照"
    assert builder._report_technique_name("firdaria", "星运-法达星限") == "法达星限"
    assert builder._report_technique_name("chart", "星盘") == "星盘"


def test_aspect_legend_only_under_wheels_that_have_aspect_lines() -> None:
    from reportlab.lib import colors
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.platypus import Paragraph

    style = ParagraphStyle("x")
    styles = {k: style for k in ("cell", "cell_muted", "cell_strong", "cell_title", "cell_big", "center_title", "center", "caption")}
    with_lines, = figures.extract_figures("chart", {"chart": _wheel_chart(), "aspects": {"normalAsp": {"Sun": {"A": [{"id": "Moon", "asp": 90}]}}}}, "星盘")
    without, = figures.extract_figures("solarreturn", {"dirChart": {"chart": _wheel_chart()}}, "太阳返照")

    def captions(fig: dict) -> str:
        blocks = figures.pdf_flowables([fig], content_width=480, font="Helvetica", bold_font="Helvetica", styles=styles,
                                       colors=colors, paragraph=Paragraph, esc=str, t=str)
        return " ".join(f.text for b in blocks for f in getattr(b, "_content", []) if isinstance(f, Paragraph))

    assert "六合" in captions(with_lines) and "六合" not in captions(without)


def test_docx_carries_the_wheel_as_a_picture(tmp_path: Path) -> None:
    import zipfile

    wheel, = figures.extract_figures("chart", {"chart": _wheel_chart()}, "星盘")
    document = {"title": "星盘", "style": "chart", "figures": [wheel],
                "sections": [{"id": "a", "title": "起盘信息", "body": "日期：2028-04-06", "human_visible": True, "group": "chart"}]}
    out = renderers.render_report(document, output_path=tmp_path / "w.docx", format_name="docx")
    assert out["format"] == "docx"
    with zipfile.ZipFile(tmp_path / "w.docx") as archive:
        media = [n for n in archive.namelist() if n.startswith("word/media/")]
        xml = archive.read("word/document.xml").decode("utf-8")
    if renderers._drawing_font_files(renderers.SIMPLIFIED) is None:
        pytest.skip("本机没有可供 Pillow 使用的 TTF 中文字体（CID 回退）")
    assert media and archive_png(tmp_path / "w.docx", media[0]).startswith(b"\x89PNG")
    assert 'w:w="11906"' in xml and 'w:h="16838"' in xml  # A4（与 PDF 同纸张）


def archive_png(path: Path, name: str) -> bytes:
    import zipfile

    with zipfile.ZipFile(path) as archive:
        return archive.read(name)


def test_sentences_after_a_diamond_stay_body_text_and_emoji_do_not_become_tofu() -> None:
    kinds = [kind for kind, _, _ in renderers._md_lines("◆ 叠断\n◆ 兄弟数：3–9 人\n◆ 互见权厄孤刃并见，刑伤已定。")]
    assert kinds == ["heading", "text", "text"]
    assert renderers._display_number("📍掌位：寅") == "▸掌位：寅"


def test_sanshiunited_does_not_redraw_its_inconsistent_qimen_sub_pan() -> None:
    cells = [{"palaceName": g, "god": "地", "tianXing": "心", "door": "杜", "tianGan": "戊", "diGan": "丙"} for g in "巽离坤震中兑艮坎乾"]
    assert figures.extract_figures("sanshiunited", {"qimen": {"cells": cells}}) == []
    assert figures.extract_figures("qimen", {"pan": {"cells": cells}})[0]["kind"] == "grid9"


def test_wheel_labels_never_overlap_even_on_a_small_crowded_wheel() -> None:
    from reportlab.graphics.shapes import String
    from reportlab.lib import colors

    chart = _wheel_chart(85.0)
    crowd = {"Sun": 89.9, "Mercury": 82.0, "Mars": 84.0, "Moon": 80.0, "Venus": 92.0, "Uranus": 87.0}  # 挤在上升点附近（盘的左侧）
    for obj in chart["objects"]:
        if obj["id"] in crowd:
            obj["lon"] = crowd[obj["id"]]
    wheel, = figures.extract_figures("chart", {"chart": chart}, "星盘")
    for size in (230.0, 380.0):  # 两图一排的小图 / 独占一行的大图
        drawing = figures.wheel_drawing(wheel, size, "Helvetica", "Helvetica-Bold", colors)
        strings = [s for s in drawing.contents if isinstance(s, String)]
        names = [s for s in strings if s.fontName == "Helvetica-Bold" and s.text in {b["label"] for b in wheel["bodies"]}]
        boxes = []
        for name in names:  # 每颗星：星名 + 正下方度数，合成一个框（宽取两行里较宽的一行，按 CJK 方字估宽）
            degree = next(s for s in strings if abs(s.x - name.x) < 0.01 and s.y < name.y and s.fontName == "Helvetica")
            half = max(len(name.text) * name.fontSize, len(degree.text) * degree.fontSize * 0.55) / 2
            boxes.append((name.x - half, degree.y - 1, name.x + half, name.y + name.fontSize * 0.8))
        for i, a in enumerate(boxes):
            for b in boxes[i + 1:]:
                assert a[2] <= b[0] or b[2] <= a[0] or a[3] <= b[1] or b[3] <= a[1], (size, a, b)


def test_key_value_keys_with_bracketed_tags_still_form_a_table() -> None:
    assert renderers._LABEL_LINE_RE.match("金 (12th; 5R12R)：1°双子54分")
    assert not renderers._LABEL_LINE_RE.match("清明后第3天；调候：无")  # 含分号的是句子，不当键
    assert not renderers._LABEL_LINE_RE.match("本报告由星阙引擎的排盘结果直接生成：收录全部盘面段落")  # 长的纯中文前缀是句子
    assert renderers._LABEL_LINE_RE.match("时间算法：真太阳时")


def test_columns_are_wide_enough_for_their_longest_number() -> None:
    widths = renderers._column_widths(["宫位", "主星", "大限"], [["夫妻宫 · 帝旺", "天府陷、文曲庙、咸池" * 3, "112~121"]], 300, 30)
    assert widths[2] >= len("112~121") * 4.6 + 11 - 0.01  # 「112~12 / 1」不再被拆开


def test_cover_shows_who_and_when_for_synastry_ingress_and_divination() -> None:
    b = builder.ReportBuilder()
    pair = b._chart_basis_section({"inner": {"name": "甲", "date": "2028-04-06", "time": "09:33:00"},
                                   "outer": {"name": "乙", "date": "1992-03-02", "time": "08:18:00"}}, {"technique_label": "关系盘"})
    assert pair["items"]["甲方"] == "甲 2028-04-06 09:33:00" and pair["items"]["乙方"].startswith("乙 1992-03-02")
    items = b._chart_basis_section({"year": 2025, "ingressTerm": "春分"}, {})["items"]
    assert items["年份"] == "2025" and items["入宫节气"] == "春分"
    assert b._chart_basis_section({"askEvent": "求财"}, {})["items"]["所问"] == "求财"


def test_docx_cover_has_no_footer_toc_has_titles_and_key_values_are_tables(tmp_path: Path) -> None:
    from docx import Document

    document = {"title": "八字 · 完整排盘报告", "style": "chart", "sections": [
        {"id": "a", "title": "起盘信息", "body": "日期：2028-04-06\n时区：+08:00\n时间算法：真太阳时", "human_visible": True, "group": "chart"},
        {"id": "b", "title": "四柱与三元", "body": "年柱：戊申", "human_visible": True, "group": "chart"},
    ]}
    renderers.render_report(document, output_path=tmp_path / "k.docx", format_name="docx")
    docx = Document(str(tmp_path / "k.docx"))
    assert docx.sections[0].different_first_page_header_footer is True
    xml = docx.element.xml
    assert "起盘信息" in xml.split("fldSimple")[1] and "四柱与三元" in xml.split("fldSimple")[1]  # 目录域的缓存结果是各章标题
    kv = [t for t in docx.tables if len(t.columns) == 2 and t.rows and t.rows[0].cells[0].text == "日期"]
    assert kv and [r.cells[0].text for r in kv[0].rows] == ["日期", "时区", "时间算法"]


def test_short_tables_are_kept_whole(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from reportlab.platypus import KeepTogether, SimpleDocTemplate

    stories: list[list] = []
    original = SimpleDocTemplate.build

    def spy(doc, flowables, *args, **kwargs):
        stories.append(list(flowables))
        return original(doc, flowables, *args, **kwargs)

    monkeypatch.setattr(SimpleDocTemplate, "build", spy)
    table = "| 步序 | 起运年 | 干支 |\n| --- | --- | --- |\n" + "\n".join(f"| 第{i}步 | {2027 + i * 10} | 甲子 |" for i in range(1, 11))
    document = {"title": "大运", "style": "chart", "sections": [{"id": "a", "title": "大运", "body": table, "human_visible": True, "group": "chart"}]}
    renderers.render_report(document, output_path=tmp_path / "t.pdf", format_name="pdf")
    # 十行的大运表整张不拆：拆开时常剩「第10步」一行孤零零在下一页顶上。
    assert any(isinstance(f, KeepTogether) for f in stories[-1])
